import json
import logging
import os
import re
from datetime import date
from io import BytesIO
from pathlib import Path

import pymupdf
from docx import Document
from dotenv import load_dotenv
from fastapi import FastAPI, File, HTTPException, UploadFile
from fastapi.responses import HTMLResponse
from openai import OpenAI, OpenAIError
from pydantic import BaseModel, Field

from tools.attendance import mark_attendance
from tools.calculator import calculate_attendance
from tools.timetable import find_day_timetable


load_dotenv()
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

MODEL_NAME = os.getenv("OPENAI_MODEL", "gpt-4o-mini")
MAX_UPLOAD_BYTES = 4 * 1024 * 1024
FRONTEND_PATH = Path(__file__).resolve().parent / "web" / "index.html"
DEFAULT_TIMETABLE_PATH = Path(__file__).resolve().parent / "test_timetable.txt"

app = FastAPI(title="College AI Assistant")


class AttendanceCalculation(BaseModel):
    attended: float = Field(ge=0)
    total: float = Field(gt=0)


class AttendanceSubmission(BaseModel):
    student_name: str = Field(min_length=1, max_length=120)
    subject: str = Field(min_length=1, max_length=120)
    attendance_date: date
    status: str


class AssistantQuestion(BaseModel):
    question: str = Field(min_length=1, max_length=4000)
    student_name: str = Field(default="Rishabh Kumar", max_length=120)
    document_text: str = Field(default="", max_length=200_000)


@app.get("/", response_class=HTMLResponse)
def home() -> HTMLResponse:
    try:
        page = FRONTEND_PATH.read_text(encoding="utf-8")
    except OSError as error:
        logger.exception("Could not load the web interface")
        raise HTTPException(
            status_code=500,
            detail="The web interface could not be loaded.",
        ) from error
    return HTMLResponse(page)


@app.get("/api/health")
def health() -> dict[str, bool | str]:
    has_gemini = bool(os.getenv("GEMINI_API_KEY", "").strip())
    has_openai = bool(os.getenv("OPENAI_API_KEY", "").strip())
    provider = "gemini" if has_gemini else ("openai" if has_openai else "none")
    return {
        "status": "ok",
        "ai_configured": has_gemini or has_openai,
        "provider": provider,
    }



@app.get("/api/default-document")
def default_document() -> dict[str, str]:
    if not DEFAULT_TIMETABLE_PATH.is_file():
        return {"name": "", "text": ""}
    try:
        text = DEFAULT_TIMETABLE_PATH.read_text(encoding="utf-8")
    except OSError as error:
        logger.exception("Could not load the default timetable")
        raise HTTPException(
            status_code=500,
            detail="The default timetable could not be loaded.",
        ) from error
    return {"name": DEFAULT_TIMETABLE_PATH.name, "text": text}


@app.post("/api/attendance/calculate")
def attendance_calculate(payload: AttendanceCalculation) -> dict[str, float]:
    try:
        percentage = calculate_attendance(payload.attended, payload.total)
    except ValueError as error:
        raise HTTPException(status_code=422, detail=str(error)) from error
    return {"percentage": percentage}


@app.post("/api/attendance/mark")
def attendance_mark(payload: AttendanceSubmission) -> dict[str, dict[str, str]]:
    try:
        record = mark_attendance(
            payload.student_name,
            payload.subject,
            payload.attendance_date.isoformat(),
            payload.status,
        )
    except ValueError as error:
        raise HTTPException(status_code=422, detail=str(error)) from error
    return {"record": record}


def extract_pdf_text_from_bytes(content: bytes) -> str:
    pages_text = []
    try:
        with pymupdf.open(stream=content, filetype="pdf") as pdf:
            for page in pdf:
                text = page.get_text("text").strip()
                if not text:
                    blocks = page.get_text("blocks")
                    text_blocks = [
                        b[4].strip()
                        for b in blocks
                        if len(b) >= 5 and isinstance(b[4], str) and b[4].strip()
                    ]
                    text = "\n".join(text_blocks)

                if not text:
                    annot_texts = []
                    for annot in page.annots() or []:
                        info = annot.info
                        if info and info.get("content"):
                            annot_texts.append(info["content"])
                    text = "\n".join(annot_texts)

                if not text:
                    try:
                        import shutil
                        import pytesseract
                        from PIL import Image

                        if shutil.which("tesseract") is not None:
                            pixmap = page.get_pixmap(dpi=150)
                            img = Image.frombytes(
                                "RGB",
                                [pixmap.width, pixmap.height],
                                pixmap.samples,
                            )
                            text = pytesseract.image_to_string(img, lang="eng").strip()
                    except Exception:
                        pass

                if text:
                    pages_text.append(text)
    except Exception as error:
        logger.exception("Error extracting PDF text")
        raise HTTPException(
            status_code=422,
            detail=f"Could not read text from this PDF file: {error}",
        ) from error

    return "\n\n".join(pages_text).strip()


def extract_document_text(filename: str, content: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    try:
        if suffix == ".pdf":
            return extract_pdf_text_from_bytes(content)
        if suffix == ".docx":
            document = Document(BytesIO(content))
            paragraphs = [p.text.strip() for p in document.paragraphs if p.text.strip()]
            for table in document.tables:
                for row in table.rows:
                    row_text = " | ".join(
                        cell.text.strip() for cell in row.cells if cell.text.strip()
                    )
                    if row_text:
                        paragraphs.append(row_text)
            return "\n".join(paragraphs)
        if suffix == ".txt":
            return content.decode("utf-8-sig")
    except HTTPException:
        raise
    except (ValueError, UnicodeDecodeError, OSError) as error:
        raise HTTPException(
            status_code=422,
            detail="Could not read this file. Check that it is a valid PDF, DOCX, or UTF-8 text file.",
        ) from error
    except Exception as error:
        logger.exception("Document extraction failed")
        raise HTTPException(
            status_code=422,
            detail="Could not read this file. Check that it is a valid PDF, DOCX, or UTF-8 text file.",
        ) from error

    raise HTTPException(
        status_code=415,
        detail="Unsupported file type. Upload a PDF, DOCX, or TXT file.",
    )



@app.post("/api/documents/upload")
async def upload_document(file: UploadFile = File(...)) -> dict[str, str]:
    filename = Path(file.filename or "").name
    if Path(filename).suffix.lower() not in {".pdf", ".docx", ".txt"}:
        raise HTTPException(
            status_code=415,
            detail="Unsupported file type. Upload a PDF, DOCX, or TXT file.",
        )
    content = await file.read(MAX_UPLOAD_BYTES + 1)
    if len(content) > MAX_UPLOAD_BYTES:
        raise HTTPException(
            status_code=413,
            detail="File is too large. The maximum upload size is 4 MB.",
        )
    text = extract_document_text(filename, content)
    return {"name": filename, "text": text}


def _client() -> tuple[OpenAI, str]:
    gemini_key = os.getenv("GEMINI_API_KEY", "").strip()
    openai_key = os.getenv("OPENAI_API_KEY", "").strip()

    if gemini_key:
        model = os.getenv("GEMINI_MODEL", "gemini-2.5-flash")
        client = OpenAI(
            api_key=gemini_key,
            base_url="https://generativelanguage.googleapis.com/v1beta/openai/",
        )
        return client, model
    elif openai_key:
        model = os.getenv("OPENAI_MODEL", "gpt-4o-mini")
        client = OpenAI(api_key=openai_key)
        return client, model
    else:
        raise HTTPException(
            status_code=503,
            detail=(
                "AI answers are not configured. Add GEMINI_API_KEY or OPENAI_API_KEY in "
                "your .env file or Vercel Environment Variables, then restart/redeploy."
            ),
        )


@app.post("/api/ask")
def ask_assistant(payload: AssistantQuestion) -> dict[str, object]:
    question = payload.question.strip()
    if not question:
        raise HTTPException(status_code=422, detail="Enter a question.")

    day_names = (
        "Monday",
        "Tuesday",
        "Wednesday",
        "Thursday",
        "Friday",
        "Saturday",
        "Sunday",
    )
    for day_name in day_names:
        if re.search(rf"\b{day_name}\b", question, re.IGNORECASE):
            if not payload.document_text.strip():
                raise HTTPException(
                    status_code=422,
                    detail="Timetable data is not available. Upload a timetable first.",
                )
            entries = find_day_timetable(payload.document_text, day_name)
            return {
                "answer": (
                    "\n".join(f"• {entry}" for entry in entries)
                    if entries
                    else f"No timetable entries were found for {day_name}."
                ),
                "tool_used": "Timetable Tool",
            }

    client, model_name = _client()
    tools = [
        {
            "type": "function",
            "function": {
                "name": "calculate_attendance",
                "description": "Calculate attendance percentage.",
                "parameters": {
                    "type": "object",
                    "properties": {
                        "attended": {"type": "number"},
                        "total": {"type": "number"},
                    },
                    "required": ["attended", "total"],
                },
            },
        },
        {
            "type": "function",
            "function": {
                "name": "mark_attendance",
                "description": "Mark student attendance.",
                "parameters": {
                    "type": "object",
                    "properties": {
                        "subject": {"type": "string"},
                        "status": {"type": "string", "enum": ["Present", "Absent"]},
                    },
                    "required": ["subject", "status"],
                },
            },
        },
    ]

    system_prompt = f"""You are College AI Assistant.
Student name: {payload.student_name or "Rishabh Kumar"}
Today's date: {date.today().isoformat()}

Use the Calculator Tool for attendance percentage.
Use the Attendance Tool when the student wants to mark attendance.
Use the supplied college document when relevant. Never invent information.
Always mention any tool used.
"""
    user_content = question
    if payload.document_text.strip():
        user_content = f"Document Context:\n{payload.document_text.strip()}\n\nQuestion:\n{question}"

    messages = [
        {"role": "system", "content": system_prompt},
        {"role": "user", "content": user_content},
    ]

    try:
        response = client.chat.completions.create(
            model=model_name,
            messages=messages,
            tools=tools,
            tool_choice="auto",
        )

        response_message = response.choices[0].message
        tools_used = []

        if response_message.tool_calls:
            messages.append(response_message)
            for tool_call in response_message.tool_calls:
                arguments = json.loads(tool_call.function.arguments)
                if tool_call.function.name == "calculate_attendance":
                    result = {
                        "percentage": calculate_attendance(
                            float(arguments["attended"]),
                            float(arguments["total"]),
                        )
                    }
                    tools_used.append("Calculator Tool")
                elif tool_call.function.name == "mark_attendance":
                    result = mark_attendance(
                        payload.student_name or "Rishabh Kumar",
                        arguments["subject"],
                        date.today().isoformat(),
                        arguments["status"],
                    )
                    tools_used.append("Attendance Tool")
                else:
                    result = {}

                messages.append(
                    {
                        "role": "tool",
                        "tool_call_id": tool_call.id,
                        "content": json.dumps(result),
                    }
                )

            final_response = client.chat.completions.create(
                model=model_name,
                messages=messages,
            )
            answer = final_response.choices[0].message.content or ""
        else:
            answer = response_message.content or ""
    except Exception as error:
        logger.exception("AI request failed")
        raise HTTPException(
            status_code=502,
            detail=f"The AI service request failed: {error}",
        ) from error

    return {"answer": answer, "tool_used": ", ".join(tools_used)}

